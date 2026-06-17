# -*- coding: utf-8 -*-
"""Вкладка «Скачать из ЕИС» для основной программы."""
from pathlib import Path

from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel, QLineEdit,
    QTextEdit, QFileDialog, QListWidget, QListWidgetItem, QCheckBox,
    QMessageBox, QGroupBox, QProgressBar,
)
from PyQt5.QtCore import Qt, QThread, pyqtSignal

from eis_downloader import EISClient, parse_numbers, download_contract
from parser_engine import ContractParser
from excel_handler import write_contracts_to_excel, create_new_excel


def read_numbers_from_excel(path):
    """Читает номера из первого столбца Excel (.xlsx)."""
    from openpyxl import load_workbook
    nums = []
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    for row in ws.iter_rows(values_only=True):
        if row and row[0] is not None:
            s = str(row[0]).strip()
            if s:
                nums.append(s)
    return nums


class TestThread(QThread):
    done = pyqtSignal(bool, str)

    def run(self):
        try:
            ok, msg = EISClient().test_connection()
        except Exception as e:
            ok, msg = False, f"Ошибка: {e}"
        self.done.emit(ok, msg)


class DownloadThread(QThread):
    progress = pyqtSignal(int, str)
    item_done = pyqtSignal(str, str, str)   # number, status, message
    finished_all = pyqtSignal(list)         # list of downloaded html paths

    def __init__(self, numbers, base_dir):
        super().__init__()
        self.numbers = numbers
        self.base_dir = base_dir

    def run(self):
        try:
            client = EISClient()
        except Exception as e:
            self.item_done.emit("", "error", f"не скачался((( ({e})")
            self.finished_all.emit([])
            return

        seen_reestr = {}
        htmls = []
        total = max(1, len(self.numbers))
        for i, num in enumerate(self.numbers):
            self.progress.emit(int(i / total * 100), f"Обработка: {num}")
            try:
                res = download_contract(client, num, self.base_dir, log=lambda m: None)
            except Exception as e:
                self.item_done.emit(num, "error", f"не скачался((( ({e})")
                continue
            reestr = res.get("reestr")
            if reestr and reestr in seen_reestr:
                self.item_done.emit(num, "skip",
                                    f"тот же контракт, что и {seen_reestr[reestr]} — пропущен")
                continue
            if reestr:
                seen_reestr[reestr] = num
            if res["status"] == "ok" and res.get("html"):
                htmls.append(res["html"])
            self.item_done.emit(num, res["status"], res["message"])
        self.progress.emit(100, "")
        self.finished_all.emit(htmls)


class EisTab(QWidget):
    def __init__(self, parent=None, on_load_to_parser=None):
        super().__init__(parent)
        self.statuses = {}        # number -> status
        self.last_htmls = []      # пути скачанных печатных форм
        self.on_load_to_parser = on_load_to_parser
        self.test_thread = None
        self.dl_thread = None
        self._build_ui()

    def _build_ui(self):
        layout = QHBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        # --- Левая панель: ввод номеров ---
        left = QGroupBox("Номера контрактов")
        lv = QVBoxLayout(left)
        hint = QLabel("Вставьте номера через запятую или столбцом "
                      "(реестровый / извещение / ИКЗ / внутренний)")
        hint.setWordWrap(True)
        hint.setStyleSheet("color:#888; font-size:11px;")
        lv.addWidget(hint)

        self.numbers_edit = QTextEdit()
        self.numbers_edit.setPlaceholderText("01085-ФЛ/2026\n2745313582726000543\n...")
        lv.addWidget(self.numbers_edit)

        self.btn_excel_load = QPushButton("Загрузить список из Excel")
        self.btn_excel_load.clicked.connect(self.load_excel_list)
        lv.addWidget(self.btn_excel_load)
        layout.addWidget(left, stretch=2)

        # --- Центр: настройки и кнопки ---
        center = QWidget()
        cv = QVBoxLayout(center)

        cv.addWidget(QLabel("Папка для скачивания:"))
        folder_row = QHBoxLayout()
        self.folder_edit = QLineEdit()
        self.folder_edit.setPlaceholderText("Куда сохранять контракты...")
        folder_row.addWidget(self.folder_edit)
        b1 = QPushButton("Обзор...")
        b1.clicked.connect(self.browse_folder)
        folder_row.addWidget(b1)
        cv.addLayout(folder_row)

        cv.addWidget(QLabel("Файл Excel (для выгрузки):"))
        excel_row = QHBoxLayout()
        self.excel_edit = QLineEdit()
        self.excel_edit.setPlaceholderText("Путь к .xlsx (по шаблону кодирование)...")
        excel_row.addWidget(self.excel_edit)
        b2 = QPushButton("Обзор...")
        b2.clicked.connect(self.browse_excel)
        excel_row.addWidget(b2)
        cv.addLayout(excel_row)

        self.chk_to_parser = QCheckBox("После скачивания загрузить в парсер")
        self.chk_to_parser.setChecked(True)
        cv.addWidget(self.chk_to_parser)
        self.chk_to_excel = QCheckBox("И сразу выгрузить из парсера в Excel")
        cv.addWidget(self.chk_to_excel)

        self.btn_test = QPushButton("ТЕСТ доступа к ЕИС")
        self.btn_test.setStyleSheet(
            "QPushButton{background:#9C27B0;color:white;border:none;border-radius:6px;"
            "padding:8px;font-weight:bold;} QPushButton:hover{background:#7B1FA2;}")
        self.btn_test.clicked.connect(self.run_test)
        cv.addWidget(self.btn_test)

        self.btn_download = QPushButton("СКАЧАТЬ")
        self.btn_download.setStyleSheet(
            "QPushButton{background:#2196F3;color:white;border:none;border-radius:8px;"
            "padding:14px;font-size:15px;font-weight:bold;} QPushButton:hover{background:#1976D2;}"
            "QPushButton:disabled{background:#bbb;}")
        self.btn_download.clicked.connect(lambda: self.start_download(False))
        cv.addWidget(self.btn_download)

        self.btn_retry = QPushButton("Скачать заново (только ❌)")
        self.btn_retry.setStyleSheet(
            "QPushButton{background:#FF9800;color:white;border:none;border-radius:6px;"
            "padding:8px;font-weight:bold;} QPushButton:hover{background:#F57C00;}")
        self.btn_retry.clicked.connect(lambda: self.start_download(True))
        cv.addWidget(self.btn_retry)

        self.progress = QProgressBar()
        self.progress.setVisible(False)
        cv.addWidget(self.progress)
        cv.addStretch()
        layout.addWidget(center, stretch=1)

        # --- Правая панель: результат и лог ---
        right = QGroupBox("Результат")
        rv = QVBoxLayout(right)
        self.result_list = QListWidget()
        rv.addWidget(self.result_list)

        res_btns = QHBoxLayout()
        self.btn_clear_res = QPushButton("Очистить")
        self.btn_clear_res.clicked.connect(self.clear_results)
        res_btns.addWidget(self.btn_clear_res)
        self.btn_to_parser = QPushButton("Загрузить в парсер")
        self.btn_to_parser.setStyleSheet(
            "QPushButton{background:#4CAF50;color:white;border:none;border-radius:6px;"
            "padding:8px;font-weight:bold;} QPushButton:hover{background:#45a049;}")
        self.btn_to_parser.clicked.connect(self.load_to_parser)
        res_btns.addWidget(self.btn_to_parser)
        rv.addLayout(res_btns)

        rv.addWidget(QLabel("Лог:"))
        self.log_area = QTextEdit()
        self.log_area.setReadOnly(True)
        self.log_area.setMaximumHeight(150)
        self.log_area.setStyleSheet("font-family:Consolas,monospace; font-size:11px;")
        rv.addWidget(self.log_area)
        layout.addWidget(right, stretch=2)

    # ------- helpers -------
    def log(self, m):
        self.log_area.append(m)

    def browse_folder(self):
        d = QFileDialog.getExistingDirectory(self, "Папка для скачивания")
        if d:
            self.folder_edit.setText(d)

    def browse_excel(self):
        path, _ = QFileDialog.getSaveFileName(self, "Файл Excel для выгрузки",
                                              "кодирование.xlsx", "Excel (*.xlsx)")
        if path:
            if not path.lower().endswith(".xlsx"):
                path += ".xlsx"
            self.excel_edit.setText(path)

    def load_excel_list(self):
        path, _ = QFileDialog.getOpenFileName(self, "Excel со списком номеров", "",
                                              "Excel (*.xlsx *.xls)")
        if not path:
            return
        try:
            nums = read_numbers_from_excel(path)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось прочитать Excel: {e}")
            return
        cur = self.numbers_edit.toPlainText().strip()
        self.numbers_edit.setPlainText((cur + "\n" if cur else "") + "\n".join(nums))
        self.log(f"Из Excel добавлено номеров: {len(nums)}")

    def run_test(self):
        self.btn_test.setEnabled(False)
        self.log("Проверка доступа к ЕИС...")
        self.test_thread = TestThread()
        self.test_thread.done.connect(self._test_done)
        self.test_thread.start()

    def _test_done(self, ok, msg):
        self.btn_test.setEnabled(True)
        self.log(msg)
        if ok:
            QMessageBox.information(self, "ТЕСТ", "Всё хорошо — ЕИС доступен.")
        else:
            QMessageBox.critical(self, "ТЕСТ", msg)

    def start_download(self, retry_only):
        base = self.folder_edit.text().strip()
        if not base or not Path(base).exists():
            QMessageBox.warning(self, "Нет папки", "Укажите существующую папку для скачивания.")
            return
        # если стоит галка выгрузки в Excel — путь обязателен заранее
        if self.chk_to_excel.isChecked() and not self.excel_edit.text().strip():
            QMessageBox.warning(self, "Не задан Excel",
                                "Укажите путь к файлу Excel (галка «выгрузить в Excel» стоит).")
            return

        if retry_only:
            numbers = [n for n, s in self.statuses.items() if s == "error"]
            if not numbers:
                QMessageBox.information(self, "Нет ошибок", "Нет номеров со статусом ❌.")
                return
        else:
            numbers = parse_numbers(self.numbers_edit.toPlainText())
            if not numbers:
                QMessageBox.warning(self, "Нет номеров", "Вставьте номера контрактов.")
                return
            seen = set()
            uniq = []
            for n in numbers:
                if n not in seen:
                    seen.add(n)
                    uniq.append(n)
            numbers = uniq
            self.result_list.clear()
            self.statuses = {}

        self.btn_download.setEnabled(False)
        self.btn_retry.setEnabled(False)
        self.progress.setVisible(True)
        self.progress.setValue(0)
        self.log(f"Старт скачивания: {len(numbers)} номер(ов)...")

        self.dl_thread = DownloadThread(numbers, base)
        self.dl_thread.progress.connect(self._on_progress)
        self.dl_thread.item_done.connect(self._on_item)
        self.dl_thread.finished_all.connect(self._on_finished)
        self.dl_thread.start()

    def _on_progress(self, pct, msg):
        self.progress.setValue(pct)

    def _on_item(self, number, status, message):
        self.statuses[number] = status
        icon = {"ok": "✅", "skip": "⏭", "error": "❌"}.get(status, "•")
        text = f"{icon} {number} — {message}"
        for i in range(self.result_list.count()):
            it = self.result_list.item(i)
            if it.data(Qt.UserRole) == number:
                it.setText(text)
                break
        else:
            it = QListWidgetItem(text)
            it.setData(Qt.UserRole, number)
            self.result_list.addItem(it)
        self.log(text)

    def clear_results(self):
        self.result_list.clear()
        self.statuses = {}
        self.last_htmls = []
        self.log_area.clear()

    def load_to_parser(self):
        if not self.last_htmls:
            QMessageBox.information(self, "Нет данных",
                                   "Сначала скачайте контракты.")
            return 0
        if not self.on_load_to_parser:
            QMessageBox.warning(self, "Недоступно", "Связка с парсером недоступна.")
            return 0
        n = self.on_load_to_parser(self.last_htmls)
        self.log(f"Загружено в парсер печатных форм: {n}")
        return n

    def _to_excel(self, htmls):
        excel_path = self.excel_edit.text().strip()
        if not excel_path:
            QMessageBox.warning(self, "Не задан Excel", "Укажите путь к файлу Excel.")
            return
        if not excel_path.lower().endswith(".xlsx"):
            excel_path += ".xlsx"
        try:
            parser = ContractParser()
            contracts = []
            for h in htmls:
                for r in parser.parse_file(h):
                    if r.contract_number and not r.is_empty:
                        contracts.append(r)
            if not contracts:
                QMessageBox.warning(self, "Excel", "Не удалось извлечь данные из скачанного.")
                return
            ok, msg = write_contracts_to_excel(excel_path, contracts)
            self.log(msg)
            if ok:
                QMessageBox.information(self, "Готово", "Выгружено в Excel:\n" + excel_path)
            else:
                QMessageBox.critical(self, "Ошибка Excel", msg)
        except Exception as e:
            QMessageBox.critical(self, "Ошибка Excel", str(e))

    def _on_finished(self, htmls):
        self.last_htmls = list(htmls)
        self.progress.setValue(100)
        self.btn_download.setEnabled(True)
        self.btn_retry.setEnabled(True)
        ok = sum(1 for s in self.statuses.values() if s == "ok")
        err = sum(1 for s in self.statuses.values() if s == "error")
        self.log(f"Готово. Успешно: {ok}, не скачалось: {err}")

        # авто-цепочка
        if self.chk_to_parser.isChecked() and self.last_htmls:
            self.load_to_parser()
        if self.chk_to_excel.isChecked() and self.last_htmls:
            self._to_excel(self.last_htmls)
        if not (self.chk_to_parser.isChecked() or self.chk_to_excel.isChecked()):
            QMessageBox.information(self, "Готово",
                                   f"Скачано: {ok}. Не скачалось: {err}.")
